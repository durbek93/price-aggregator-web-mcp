# -*- coding: utf-8 -*-
"""Веб-версия конвейера (pipeline) для обработки прайс-листов.
Принимает конфигурации из базы данных (в виде словарей и списков) и запускает фоновые задачи.
"""
from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path
from typing import Callable, List, Optional, Tuple, Dict, Any

import openpyxl
from openpyxl.styles import PatternFill, Font
import pandas as pd

# Динамически добавляем путь к текущей папке в sys.path, чтобы импорты внутри модулей работали
core_dir = str(Path(__file__).parent.absolute())
if core_dir not in sys.path:
    sys.path.insert(0, core_dir)

import assortment_prices as assortment
import best_suppliers_analysis as best_suppliers
import cheapest_svodny_prices as cheapest
from svodny_price import (
    build_matrix_from_extracted,
    extract_materia_any,
    format_summary_sheet,
    friendly_source_label,
    fuzzy_merge_canonicals,
    read_price_table,
    is_excluded_product,
    match_canonical,
    is_materia_producer_norm,
    pick_price_column,
    pick_column,
    load_distributor_rules,
)
from price_utils import to_number, round_price, normalize_name
import shutil
from rapidfuzz import fuzz
from svodny_price_vertical import (
    build_matrix_from_extracted as build_vertical_matrix,
    format_summary_sheet as format_vertical_sheet,
)


def _column_lists(settings_dict: Dict[str, Any]) -> Tuple[list, list, list, list, list]:
    c = settings_dict.get("columns", {})
    prod_c = list(c.get("producer_candidates", []))
    for extra in ["Производитель", "Бренд", "Завод", "Фирма", "Произ.", "Произв.", "Произв", "Произ-ль", "Произв-ль", "Manufacturer", "Производство"]:
        if extra not in prod_c:
            prod_c.append(extra)
    return (
        list(c.get("key_candidates", [])),
        list(c.get("name_candidates", [])),
        list(c.get("price_candidates", [])),
        list(c.get("stock_candidates", [])),
        prod_c,
    )


def _display_label(entry_display: str, filename: str) -> str:
    d = (entry_display or "").strip()
    if d:
        return d
    return friendly_source_label(Path(filename).stem)


def _extract_matrix_rows(
    settings_dict: Dict[str, Any],
    catalog_list: List[Tuple[str, List[str]]],
    price_files: List[Dict[str, Any]],
    upload_dir: Path,
    progress: Optional[Callable[[str], None]] = None,
) -> Tuple[List[dict], List[Tuple[str, List[str]]], List[str]]:
    
    use_catalog_matching = settings_dict.get("use_catalog_matching", False)
    producers = settings_dict.get("producers", [])
    legacy_priority_rules = settings_dict.get("legacy_priority_rules", False)

    if use_catalog_matching and not catalog_list:
        raise ValueError("Добавьте препараты в настройках (каталог: название и синонимы).")
    if not producers:
        raise ValueError("Добавьте хотя бы одну подстроку производителя в настройках.")

    key_c, name_c, price_c, stock_c, prod_c = _column_lists(settings_dict)
    
    # Filter files that actually exist
    valid_entries = []
    for e in price_files:
        filepath = upload_dir / e["filename"]
        if filepath.is_file():
            valid_entries.append(e)

    if not valid_entries:
        raise ValueError("Нет файлов прайсов на сервере. Пожалуйста, загрузите прайсы.")

    all_labels: List[str] = []
    extracted: List[dict] = []
    total = len(valid_entries)
    for i, e in enumerate(valid_entries, 1):
        path = upload_dir / e["filename"]
        if progress:
            progress(f"Обработка [{i}/{total}]: {path.name}")
        preset = (e.get("read_preset") or "").strip() or None
        try:
            df = read_price_table(
                path,
                key_c,
                name_c,
                price_c,
                stock_c,
                use_key=True,
                use_name=True,
                read_preset=preset,
            )
            lbl = _display_label(e.get("display_name", ""), e["filename"])
            items = extract_materia_any(
                df,
                path,
                key_c,
                name_c,
                price_c,
                stock_c,
                prod_c,
                catalog_list if use_catalog_matching else [],
                source_label=lbl,
                producer_substrings=producers,
                legacy_priority_rules=legacy_priority_rules,
            )
            extracted.extend(items)
            all_labels.append(lbl)
        except Exception as ex:
            import logging
            logging.getLogger(__name__).exception(f"Error processing file {path.name}: {ex}")
            if progress:
                progress(f"⚠️ Пропущен: Ошибка в файле {path.name}: {str(ex)}")

    if not all_labels:
        raise ValueError("Ни один из файлов прайс-листов не удалось обработать.")

    all_labels = sorted(dict.fromkeys(all_labels))
    return fuzzy_merge_canonicals(extracted, threshold=82, catalog=catalog_list if use_catalog_matching else None), (catalog_list if use_catalog_matching else []), all_labels


def _save_matrix(df: pd.DataFrame, out: Path, vertical: bool) -> Path:
    out.parent.mkdir(parents=True, exist_ok=True)
    df.to_excel(out, index=False, engine="openpyxl")
    wb = openpyxl.load_workbook(out)
    for sheet in wb.worksheets:
        if vertical:
            format_vertical_sheet(sheet)
        else:
            format_summary_sheet(sheet)
    wb.save(out)
    return out


def pick_price_columns(columns: List[str], candidates: List[str]) -> Tuple[Optional[str], Optional[str]]:
    local_candidates = list(candidates)
    for extra in ["Продажная", "Продажная цена", "Продажная с НДС"]:
        if extra not in local_candidates:
            local_candidates.append(extra)
            
    def _super_normalize(col: str) -> str:
        s = str(col).lower()
        s = s.replace("*", "").replace("®", "").replace("™", "")
        return "".join(s.split())
        
    matched_candidates = []
    for col in columns:
        if not col or not str(col).strip():
            continue
        col_norm = _super_normalize(col)
        if not col_norm:
            continue
            
        low = str(col).lower()
        # Exclude known non-selling price columns and tax-only/order-only columns
        if any(x in low for x in ["референт", "лимит", "реестр", "предел", "max", "min", "нац", "наценка", "нагрузка"]):
            continue
        if "%" in low and not any(x in low for x in ["цена", "price", "нарх"]):
            continue
        # Exclude order count, sum, tax columns
        if low in ["ндс", "ндс%", "ставка ндс", "ставка ндс, %", "ндс, %", "заказ", "заказы", "кол-во", "кол во", "сумма", "сумма заказа", "сумма заказ", "сумма со скидкой", "sum", "total", "кол-во заказ", "количество"]:
            continue
        if any(x in low for x in ["сумма", "sum", "total"]):
            continue
            
        for cand in local_candidates:
            cand_norm = _super_normalize(cand)
            if not cand_norm:
                continue
            if cand_norm == col_norm or cand_norm in col_norm:
                matched_candidates.append(col)
                break
                
    if not matched_candidates:
        return None, None
        
    matched_candidates = list(dict.fromkeys(matched_candidates))
    
    def get_score(col_name: str) -> float:
        low = col_name.lower()
        score = 0.0
        if "цена" in low or "price" in low or "нарх" in low or "продаж" in low:
            score += 10.0
        if "сумма" in low or "sum" in low:
            score -= 5.0
            
        if "отпускная" in low or "отп" in low:
            score += 5.0
        if "клиент" in low or "предоплат" in low:
            score += 4.0
        if "скид" in low or "disc" in low:
            score += 3.0
        if "с ндс" in low or "ндс" in low:
            score += 2.0
            
        if "базовая" in low or "база" in low or "баз" in low:
            score -= 4.0
        if "розница" in low or "розн" in low:
            score -= 3.0
        if "закуп" in low or "вход" in low:
            score -= 3.0
        if "заказ" in low:
            score -= 6.0
            
        return score

    matched_candidates = sorted(matched_candidates, key=get_score, reverse=True)
    
    col_25 = None
    col_100 = None
    
    for col in matched_candidates:
        col_norm = _super_normalize(col)
        if "25" in col_norm and col_25 is None:
            col_25 = col
        elif "100" in col_norm and col_100 is None:
            col_100 = col
            
    remaining = [c for c in matched_candidates if c != col_25 and c != col_100]
    
    if col_100 is not None and col_25 is None:
        valid_remaining = [c for c in remaining if get_score(c) >= 5.0]
        if valid_remaining:
            col_25 = valid_remaining[0]
        else:
            col_25 = col_100
            
    elif col_25 is not None and col_100 is None:
        valid_remaining = [c for c in remaining if get_score(c) >= 5.0]
        if valid_remaining:
            col_100 = valid_remaining[0]
        else:
            col_100 = col_25
            
    elif col_25 is None and col_100 is None:
        valid_candidates = [c for c in matched_candidates if get_score(c) >= 5.0]
        if len(valid_candidates) >= 2:
            col_100 = valid_candidates[0]
            col_25 = valid_candidates[1]
        elif len(valid_candidates) == 1:
            col_100 = valid_candidates[0]
            col_25 = valid_candidates[0]
    return col_25, col_100


def _find_matching_col(target: Optional[str], columns: List[Any]) -> Optional[str]:
    """Гибкий поиск колонки по имени (с нормализацией пробелов, регистра и .1 <-> _1)."""
    if not target:
        return None
    str_cols = [str(c) for c in columns]
    if target in str_cols:
        return target
    t_clean = str(target).strip().lower().replace(".", "_").replace(" ", "")
    for col in str_cols:
        c_clean = col.strip().lower().replace(".", "_").replace(" ", "")
        if c_clean == t_clean:
            return col
    return None


def find_distributor_columns(label: str) -> Optional[Tuple[int, Optional[int]]]:
    """Возвращает 1-based колонки (col_25, col_100) из шаблона Шаблон.xlsx."""
    lbl = label.lower()
    if "zenta" in lbl or "зента" in lbl:
        return (3, 4)
    elif "grand" in lbl or "гранд" in lbl:
        return (5, 6)
    elif "meros" in lbl or "мерос" in lbl:
        return (7, 8)
    elif "promotion" in lbl or "промоушен" in lbl or "промоушн" in lbl:
        return (9, 10)
    elif "choice" in lbl or "чойс" in lbl:
        return (11, None)
    elif "biotek" in lbl or "биотек" in lbl or "biotech" in lbl or "спец" in lbl:
        return (12, 13)
    elif "luxe" in lbl or "люкс" in lbl:
        return (14, 15)
    elif "aero" in lbl or "аэро" in lbl:
        return (16, 17)
    elif "delpharm" in lbl or "дельфарм" in lbl:
        return (18, None)
    return None


def match_product_to_template(product_name: str, template_products: List[Tuple[int, str]], original_name: Optional[str] = None) -> Optional[int]:
    """Сопоставляет название препарата со строкой в шаблоне."""
    p_norm = str(product_name).lower()
    orig_norm = str(original_name or "").lower()
    
    # Special exact mapping rules for specific brands
    p_norm_clean = " ".join(p_norm.split())
    orig_norm_clean = " ".join(orig_norm.split())
    
    def contains_any(word: str, substrings: List[str]) -> bool:
        return any(x in word for x in substrings)

    # 0. Ergoferon
    if "эргоферон" in p_norm_clean or "эргоферон" in orig_norm_clean or "ergoferon" in p_norm_clean or "ergoferon" in orig_norm_clean:
        for row_idx, t_name in template_products:
            t_low = t_name.lower()
            if "эргоферон" in t_low or "ergoferon" in t_low:
                return row_idx

    # 1. Anaferon
    if "анаферон" in p_norm_clean or "анаферон" in orig_norm_clean:
        is_drops = contains_any(p_norm_clean, ["капл", "25мл", "25 мл"]) or contains_any(orig_norm_clean, ["капл", "25мл", "25 мл"])
        is_child = contains_any(p_norm_clean, ["детс", "дет"]) or contains_any(orig_norm_clean, ["детс", "дет"])
        
        if is_drops:
            for row_idx, t_name in template_products:
                t_low = t_name.lower()
                if "анаферон" in t_low and ("капл" in t_low or "кап" in t_low):
                    return row_idx
        elif is_child:
            for row_idx, t_name in template_products:
                t_low = t_name.lower()
                if "анаферон" in t_low and "дет" in t_low and "кап" not in t_low:
                    return row_idx
        else:
            for row_idx, t_name in template_products:
                t_low = t_name.lower()
                if "анаферон" in t_low and "дет" not in t_low and "кап" not in t_low:
                    return row_idx

    # 2. Tenoten
    if "тенотен" in p_norm_clean or "тенотен" in orig_norm_clean:
        is_child = contains_any(p_norm_clean, ["детс", "дет"]) or contains_any(orig_norm_clean, ["детс", "дет"])
        if is_child:
            for row_idx, t_name in template_products:
                t_low = t_name.lower()
                if "тенотен" in t_low and "дет" in t_low:
                    return row_idx
        else:
            for row_idx, t_name in template_products:
                t_low = t_name.lower()
                if "тенотен" in t_low and "дет" not in t_low:
                    return row_idx

    # 3. Rengalin
    if "ренгалин" in p_norm_clean or "ренгалин" in orig_norm_clean:
        is_sol = contains_any(p_norm_clean, ["раствор", "р-р", "100мл", "100 мл", "сироп"]) or contains_any(orig_norm_clean, ["раствор", "р-р", "100мл", "100 мл", "сироп"])
        if is_sol:
            for row_idx, t_name in template_products:
                t_low = t_name.lower()
                if "ренгалин" in t_low and ("раствор" in t_low or "р-р" in t_low or "100" in t_low or "мл" in t_low):
                    return row_idx
        else:
            for row_idx, t_name in template_products:
                t_low = t_name.lower()
                if "ренгалин" in t_low and not ("раствор" in t_low or "р-р" in t_low or "100" in t_low or "мл" in t_low):
                    return row_idx

    # 4. Klimaksan
    if "климаксан" in p_norm_clean or "климаксан" in orig_norm_clean:
        is_gran = contains_any(p_norm_clean, ["гран", "гранул", "10г", "10 г"]) or contains_any(orig_norm_clean, ["гран", "гранул", "10г", "10 г"])
        if is_gran:
            for row_idx, t_name in template_products:
                t_low = t_name.lower()
                if "климаксан" in t_low and ("гран" in t_low or "10" in t_low):
                    return row_idx
        else:
            for row_idx, t_name in template_products:
                t_low = t_name.lower()
                if "климаксан" in t_low and not ("гран" in t_low or "10" in t_low):
                    return row_idx

    # 5. Divaza
    if "диваза" in p_norm_clean or "диваза" in orig_norm_clean:
        for row_idx, t_name in template_products:
            t_low = t_name.lower()
            if "диваза" in t_low:
                return row_idx

    # 6. Kolofort
    if "колофорт" in p_norm_clean or "колофорт" in orig_norm_clean:
        for row_idx, t_name in template_products:
            t_low = t_name.lower()
            if "колофорт" in t_low:
                return row_idx

    # 7. Afalaza
    if "афалаза" in p_norm_clean or "афалаза" in orig_norm_clean:
        for row_idx, t_name in template_products:
            t_low = t_name.lower()
            if "афалаза" in t_low:
                return row_idx

    # 8. Afala
    if ("афала" in p_norm_clean or "афала" in orig_norm_clean) and "афалаза" not in p_norm_clean and "афалаза" not in orig_norm_clean:
        for row_idx, t_name in template_products:
            t_low = t_name.lower()
            if "афала" in t_low and "афалаза" not in t_low:
                return row_idx

    # 9. Impaza
    if "импаза" in p_norm_clean or "импаза" in orig_norm_clean:
        for row_idx, t_name in template_products:
            t_low = t_name.lower()
            if "импаза" in t_low:
                return row_idx

    # 10. Mlekoin
    if "млекоин" in p_norm_clean or "млекоин" in orig_norm_clean:
        for row_idx, t_name in template_products:
            t_low = t_name.lower()
            if "млекоин" in t_low:
                return row_idx

    # 11. Uspokoy
    if "успокой" in p_norm_clean or "успокой" in orig_norm_clean:
        for row_idx, t_name in template_products:
            t_low = t_name.lower()
            if "успокой" in t_low:
                return row_idx

    # 12. Faringomed
    if "фарингомед" in p_norm_clean or "фарингомед" in orig_norm_clean:
        for row_idx, t_name in template_products:
            t_low = t_name.lower()
            if "фарингомед" in t_low:
                return row_idx

    is_drops = "капл" in p_norm
    is_sol = "раствор" in p_norm or "р-р" in p_norm or "мл" in p_norm
    
    best_row = None
    best_score = 0.0
    for row_idx, t_name in template_products:
        t_norm = t_name.lower()
        if is_drops and "капл" not in t_norm:
            continue
        if not is_drops and "капл" in t_norm:
            continue
        if is_sol and "раствор" not in t_norm and "мл" not in t_norm:
            continue
        if not is_sol and ("раствор" in t_norm or "мл" in t_norm):
            continue
            
        score = fuzz.token_sort_ratio(p_norm, t_norm)
        if score > best_score:
            best_score = score
            best_row = row_idx
            
    if best_score >= 75:
        return best_row
    return None


FILL_GREEN = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
FONT_GREEN = Font(color="375623", bold=True)

FILL_RED = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
FONT_RED = Font(color="C65911", bold=True)


def _load_previous_svodny_prices(output_dir: Path, exclude_filename: str = "") -> Dict[Tuple[int, int], float]:
    """
    Загружает цены из последнего ранее сгенерированного файла «Сводный прайс *.xlsx».
    Возвращает словарь (row, col) -> float_price.
    """
    prev_prices: Dict[Tuple[int, int], float] = {}
    if not output_dir.exists():
        return prev_prices

    candidates = []
    for p in output_dir.glob("Сводный прайс *.xlsx"):
        if p.name == exclude_filename or "вертикально" in p.name.lower():
            continue
        try:
            candidates.append((p.stat().st_mtime, p))
        except OSError:
            pass

    if not candidates:
        return prev_prices

    candidates.sort(key=lambda x: x[0], reverse=True)
    latest_file = candidates[0][1]

    try:
        wb_prev = openpyxl.load_workbook(latest_file, data_only=True)
        sheet_prev = wb_prev["СвПрайс"] if "СвПрайс" in wb_prev.sheetnames else wb_prev.active
        for r in range(7, 25):
            for c in range(3, 19):
                val = sheet_prev.cell(row=r, column=c).value
                if val is not None:
                    try:
                        if isinstance(val, (int, float)):
                            prev_prices[(r, c)] = float(val)
                        else:
                            clean_s = "".join(ch for ch in str(val) if ch.isdigit() or ch in ".,")
                            if clean_s:
                                prev_prices[(r, c)] = float(clean_s.replace(",", "."))
                    except (ValueError, TypeError):
                        pass
        wb_prev.close()
    except Exception:
        pass

    return prev_prices


def _get_current_numeric_price(cell_val: Any) -> Optional[float]:
    if cell_val is None:
        return None
    if isinstance(cell_val, (int, float)):
        return float(cell_val)
    try:
        clean_s = "".join(ch for ch in str(cell_val) if ch.isdigit() or ch in ".,")
        if clean_s:
            return float(clean_s.replace(",", "."))
    except (ValueError, TypeError):
        pass
    return None


def _set_price_cell(sheet, row: int, col: int, price_to_put: float, prev_prices: Dict[Tuple[int, int], float]):
    curr_raw = sheet.cell(row=row, column=col).value
    curr_num = _get_current_numeric_price(curr_raw)
    
    # Если цена уже установлена, перезаписываем только если новая цена выгоднее (ниже)
    if curr_num is not None and price_to_put >= curr_num:
        return

    rounded_price = round_price(price_to_put)
    prev_price = prev_prices.get((row, col))
    
    cell = sheet.cell(row=row, column=col)
    
    if prev_price is not None and prev_price > 0:
        prev_rounded = round_price(prev_price)
        if rounded_price > prev_rounded:
            val_str = f"{int(rounded_price):,}".replace(",", " ")
            cell.value = f"{val_str} ↑"
            cell.fill = FILL_RED
            cell.font = FONT_RED
            return
        elif rounded_price < prev_rounded:
            val_str = f"{int(rounded_price):,}".replace(",", " ")
            cell.value = f"{val_str} ↓"
            cell.fill = FILL_GREEN
            cell.font = FONT_GREEN
            return

    cell.value = rounded_price


def run_svodny_horizontal(
    settings_dict: Dict[str, Any],
    catalog_list: List[Tuple[str, List[str]]],
    price_files: List[Dict[str, Any]],
    upload_dir: Path,
    output_dir: Path,
    progress: Optional[Callable[[str], None]] = None,
) -> Path:
    template_path = Path("/app/Шаблон.xlsx")
    if not template_path.exists():
        template_path = Path(__file__).resolve().parents[2] / "Шаблон.xlsx"
    if not template_path.exists():
        template_path = Path(__file__).resolve().parents[1] / "Шаблон.xlsx"
        
    if not template_path.exists():
        raise FileNotFoundError(f"Шаблон.xlsx не найден: {template_path}")
        
    today = datetime.now().strftime("%d.%m.%Y")
    out = output_dir / f"Сводный прайс {today}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)

    prev_prices = _load_previous_svodny_prices(output_dir, exclude_filename=out.name)

    shutil.copyfile(template_path, out)
    
    wb = openpyxl.load_workbook(out)
    sheet = wb["СвПрайс"]
    
    # Считываем препараты шаблона (строки 7-24)
    template_products = []
    for r in range(7, 25):
        p_name = sheet.cell(row=r, column=2).value
        if p_name:
            template_products.append((r, str(p_name).strip()))
            
    # Полностью очищаем старые цены из шаблона
    for r in range(7, 25):
        for c in range(3, 19):
            sheet.cell(row=r, column=c).value = None
            
    key_c, name_c, price_c, stock_c, prod_c = _column_lists(settings_dict)
    # Гарантируем поддержку сокращенных названий колонки производителя (например, Произ.)
    for extra in ["произ.", "производ", "производ.", "произв.", "произв", "произ-ль", "произв-ль"]:
        if extra not in prod_c:
            prod_c.append(extra)
    producers = settings_dict.get("producers", [])
    legacy_priority_rules = settings_dict.get("legacy_priority_rules", False)
    use_catalog_matching = settings_dict.get("use_catalog_matching", False)
    
    valid_entries = []
    for e in price_files:
        filepath = upload_dir / e["filename"]
        if filepath.is_file():
            valid_entries.append(e)

    # Автоматическая дедупликация: если для одного дистрибьютора есть несколько файлов, берем только самый свежий
    distributor_files: Dict[str, Dict[str, Any]] = {}
    other_entries: List[Dict[str, Any]] = []

    for e in valid_entries:
        filepath = upload_dir / e["filename"]
        lbl = _display_label(e.get("display_name", ""), e["filename"])
        col_mapping = find_distributor_columns(lbl)
        if not col_mapping:
            other_entries.append(e)
            continue

        fn_low = filepath.name.lower()
        lbl_low = lbl.lower()
        dist_key = f"{col_mapping[0]}_{col_mapping[1]}"
        if any(d in lbl_low or d in fn_low for d in ["grand", "гранд", "promotion", "промоушен", "промоушн", "luxe", "люкс"]):
            variant = "100" if "100" in fn_low else "25"
            dist_key = f"{dist_key}_{variant}"

        mtime = filepath.stat().st_mtime if filepath.exists() else 0
        if dist_key in distributor_files:
            existing_path = upload_dir / distributor_files[dist_key]["filename"]
            existing_mtime = existing_path.stat().st_mtime if existing_path.exists() else 0
            if mtime > existing_mtime:
                if progress:
                    progress(f"ℹ️ Для дистрибьютора '{lbl}' выбран более свежий файл '{filepath.name}', устаревший дубликат '{distributor_files[dist_key]['filename']}' пропущен")
                distributor_files[dist_key] = e
            else:
                if progress:
                    progress(f"ℹ️ Для дистрибьютора '{lbl}' пропущен устаревший дубликат '{filepath.name}' (используется более свежий '{distributor_files[dist_key]['filename']}')")
        else:
            distributor_files[dist_key] = e

    deduped_entries = list(distributor_files.values()) + other_entries
            
    total = len(deduped_entries)
    for i, e in enumerate(deduped_entries, 1):
        path = upload_dir / e["filename"]
        if progress:
            progress(f"Обработка [{i}/{total}]: {path.name}")
            
        preset = (e.get("read_preset") or "").strip() or None
        lbl = _display_label(e.get("display_name", ""), e["filename"])
        
        col_mapping = find_distributor_columns(lbl)
        if not col_mapping:
            if progress:
                progress(f"⚠️ Дистрибьютор '{lbl}' не сопоставлен с колонками шаблона. Пропускаем.")
            continue
            
        col_25, col_100 = col_mapping
        
        try:
            df_file = read_price_table(
                path,
                key_c,
                name_c,
                price_c,
                stock_c,
                use_key=True,
                use_name=True,
                read_preset=preset,
            )
            
            name_col = pick_column(list(df_file.columns), name_c)
            producer_col = pick_column(list(df_file.columns), prod_c)
            p25_col, p100_col = pick_price_columns(list(df_file.columns), price_c)
            
            fn_low = path.name.lower()
            lbl_low = lbl.lower()
            
            # Check for custom overrides from debugger agent rules for single-file distributors
            dist_rules = load_distributor_rules()
            rule = None
            is_multi_file_dist = any(d in lbl_low or d in fn_low for d in ["grand", "гранд", "promotion", "промоушен", "промоушн", "luxe", "люкс"])
            
            if not is_multi_file_dist:
                for r_key, r_val in dist_rules.items():
                    if r_key.lower() in lbl_low or r_key.lower() in fn_low:
                        rule = r_val
                        break
            
            if rule:
                matched_name = _find_matching_col(rule.get("name_col"), list(df_file.columns))
                if matched_name:
                    name_col = matched_name
                matched_p25 = _find_matching_col(rule.get("p25_col"), list(df_file.columns))
                if matched_p25:
                    p25_col = matched_p25
                matched_p100 = _find_matching_col(rule.get("p100_col"), list(df_file.columns))
                if matched_p100:
                    p100_col = matched_p100
            elif "meros" in fn_low or "мерос" in fn_low or "meros" in lbl_low or "мерос" in lbl_low:
                # For Meros Pharm, find column for 25% and 100% manually
                for col in df_file.columns:
                    c_low = str(col).lower().strip()
                    if c_low in ["цена договорная", "ценадоговорная", "цена"]:
                        p25_col = col
                        break
                if not p25_col and "Цена" in df_file.columns:
                    p25_col = "Цена"
                    
                for col in df_file.columns:
                    c_low = str(col).lower().strip()
                    if c_low in ["цена 100%", "цена100%", "100% оплата", "100%оплата", "при 100% оплате", "при100%оплате"]:
                        p100_col = col
                        break
                if not p100_col and "Цена 100%" in df_file.columns:
                    p100_col = "Цена 100%"
            elif any(d in fn_low or d in lbl_low for d in ["grand", "гранд", "promotion", "промоушен", "промоушн"]):
                # Grand Pharm and Drug Promotion have separate 100% and 25% files
                is_100 = "100" in fn_low
                is_25 = "25" in fn_low or not is_100

                # For Grand Pharm: if file has Отп. цена без НДС and ставка НДС, compute Отпускная цена с НДС
                if any(d in fn_low or d in lbl_low for d in ["grand", "гранд"]):
                    if any("отп" in str(c).lower() and "без ндс" in str(c).lower() for c in df_file.columns) and not any("отпускная цена с ндс" in str(c).lower() for c in df_file.columns):
                        base_col = next((c for c in df_file.columns if "отп" in str(c).lower() and "без ндс" in str(c).lower()), None)
                        nds_col = next((c for c in df_file.columns if "ставка" in str(c).lower() and "ндс" in str(c).lower()), None)
                        if base_col:
                            def _calc_grand_nds(row):
                                b = to_number(row.get(base_col))
                                if b is None:
                                    return None
                                n = to_number(row.get(nds_col)) if nds_col else 0.12
                                if n is None or n <= 0:
                                    n = 0.12
                                elif n > 1.0:
                                    n = n / 100.0
                                return round(b * (1.0 + n), 2)
                            df_file["Отпускная цена с НДС"] = df_file.apply(_calc_grand_nds, axis=1)

                col_price = None
                if is_100:
                    for col in df_file.columns:
                        c_low = str(col).lower().strip()
                        if "100" in c_low and ("цена" in c_low or "%" in c_low):
                            col_price = col
                            break

                if not col_price:
                    patterns = [
                        ["отпускная цена с ндс", "отпуск.цена с ндс", "отп. цена с ндс"],
                        ["цена с ндс", "цена (с ндс)", "цена с ндс 100%", "цена с ндс 25%"],
                        ["отпускная цена", "отпуск.цена", "отп. цена", "отп.цена"],
                        ["цена", "price", "нарх"],
                    ]
                    for pattern_group in patterns:
                        for col in df_file.columns:
                            col_clean = " ".join(str(col).lower().split())
                            if "без ндс" in col_clean or "безндс" in col_clean:
                                continue
                            if "базов" in col_clean:
                                continue
                            if col_clean in ["ндс", "ндс%", "ставка ндс", "ставка ндс, %", "ндс, %", "заказ", "заказы", "кол-во", "кол во", "сумма", "sum", "total", "количество"]:
                                continue
                            if any(x in col_clean for x in ["сумма", "sum", "total"]):
                                continue
                            if any(p == col_clean or p in col_clean for p in pattern_group):
                                col_price = col
                                break
                        if col_price:
                            break
                            
                if not col_price:
                    valid_cols = [c for c in df_file.columns if "без ндс" not in str(c).lower() and "безндс" not in str(c).lower() and "базов" not in str(c).lower()]
                    col_price = pick_price_column(valid_cols, price_c)
                    
                if is_100:
                    p100_col = col_price
                    p25_col = None
                else:
                    p25_col = col_price
                    p100_col = None
            elif "luxe" in fn_low or "люкс" in fn_low or "luxe" in lbl_low or "люкс" in lbl_low:
                # For Pharm Luxe, separate 100% and 25% files
                col_price = None
                patterns = [
                    ["цена с ндс", "цена (с ндс)", "цена с ндс 100%", "цена с ндс 25%"],
                    ["отпускная цена с ндс", "отпуск.цена с ндс", "отп. цена с ндс"],
                    ["отпускная цена", "отпуск.цена", "отп. цена", "отп.цена"],
                    ["цена", "price", "нарх"],
                ]
                for pattern_group in patterns:
                    for col in df_file.columns:
                        col_clean = " ".join(str(col).lower().split())
                        if "без ндс" in col_clean or "безндс" in col_clean:
                            continue
                        if "базов" in col_clean:
                            continue
                        if col_clean in ["ндс", "ндс%", "ставка ндс", "ставка ндс, %", "ндс, %", "заказ", "заказы", "кол-во", "кол во", "сумма", "sum", "total", "количество"]:
                            continue
                        if any(x in col_clean for x in ["сумма", "sum", "total"]):
                            continue
                        if any(p == col_clean or p in col_clean for p in pattern_group):
                            col_price = col
                            break
                    if col_price:
                        break
                        
                if not col_price:
                    valid_cols = [c for c in df_file.columns if "без ндс" not in str(c).lower() and "безндс" not in str(c).lower() and "базов" not in str(c).lower()]
                    col_price = pick_price_column(valid_cols, price_c)
                    
                if "100" in fn_low:
                    p100_col = col_price
                    p25_col = None
                else:
                    p25_col = col_price
                    p100_col = None
            elif "biotek" in fn_low or "биотек" in fn_low or "biotek" in lbl_low or "биотек" in lbl_low or "biotech" in fn_low or "спец" in fn_low or "спец" in lbl_low:
                # For Biotek Pharm
                raw_cols = list(df_file.columns)
                name_col_idx = raw_cols.index(name_col) if name_col in raw_cols else None

                # Look for explicit discount / special price columns
                explicit_100_col = None
                explicit_25_col = None
                for col in raw_cols:
                    c_clean = str(col).strip().lower().replace("–", "-").replace("—", "-").replace("−", "-")
                    if any(c_clean == t or c_clean.replace(" ", "") == t.replace(" ", "") or t in c_clean for t in [
                        "цена-4", "цена 4", "цена - 4", "-4", "-4%", 
                        "цена-3", "цена 3", "цена - 3", "-3", "-3%", 
                        "цена-2", "цена 2", "цена - 2", "-2", "-2%", 
                        "100%", "цена 100%", "спец прайс", "спец. прайс", "спецпрайс",
                        "спец цена", "спец. цена", "спеццена", "цена спец", "цена со скидкой"
                    ]):
                        explicit_100_col = col
                    if any(c_clean == t or c_clean.replace(" ", "") == t.replace(" ", "") for t in ["цена", "базовая цена", "отпускная цена", "отп.цена", "цена без скидки"]):
                        if not explicit_25_col:
                            explicit_25_col = col

                min_valid = min(10, max(2, int(len(df_file) * 0.1)))
                price_col_indices = []
                for idx, col in enumerate(raw_cols):
                    if name_col_idx is not None and idx <= name_col_idx:
                        continue
                    c_str = str(col).lower().replace("–", "-").replace("—", "-").replace("−", "-").strip()
                    if any(x in c_str for x in ["заказ", "орг. уп", "орг.уп", "произ.", "производитель", "годности", "годность", "остаток", "наличие"]):
                        break

                    valid_prices = 0
                    for v in df_file.iloc[:, idx].dropna():
                        try:
                            if float(v) > 500:
                                valid_prices += 1
                        except (ValueError, TypeError):
                            pass
                    if valid_prices >= min_valid:
                        price_col_indices.append(idx)

                if explicit_25_col and explicit_100_col:
                    p25_col = explicit_25_col
                    p100_col = explicit_100_col
                elif explicit_100_col and price_col_indices:
                    p25_col = raw_cols[price_col_indices[0]]
                    p100_col = explicit_100_col
                elif len(price_col_indices) >= 2:
                    p25_col = raw_cols[price_col_indices[0]]
                    p100_col = raw_cols[price_col_indices[-1]]
                elif price_col_indices:
                    p25_col = raw_cols[price_col_indices[0]]
                    p100_col = raw_cols[price_col_indices[0]]
                else:
                    p25_col = pick_price_column(raw_cols, price_c)
                    p100_col = p25_col

            
            if not name_col:
                raise ValueError("Не найдена колонка наименования препарата.")
                
            for _, r_data in df_file.iterrows():
                name_val = r_data.get(name_col)
                if name_val is None or (isinstance(name_val, float) and pd.isna(name_val)):
                    continue
                if is_excluded_product(str(name_val)):
                    continue
                    
                producer_val = r_data.get(producer_col) if producer_col else ""
                prod_norm = normalize_name(str(producer_val or ""))
                
                name_norm = normalize_name(str(name_val))
                
                canonical = match_canonical(
                    name_norm, catalog_list if use_catalog_matching else [], legacy_priority_rules=legacy_priority_rules
                )
                
                by_producer = False
                if producers:
                    subs = [normalize_name(s) for s in producers if str(s).strip()]
                    by_producer = any(s and s in prod_norm for s in subs)
                else:
                    by_producer = is_materia_producer_norm(prod_norm)
                    
                if not by_producer and not canonical:
                    continue
                    
                prod_name_to_match = canonical or str(name_val)
                matched_row = match_product_to_template(prod_name_to_match, template_products, str(name_val))
                if not matched_row:
                    continue
                    
                val_25 = to_number(r_data.get(p25_col)) if p25_col else None
                val_100 = to_number(r_data.get(p100_col)) if p100_col else None
                
                if col_100 is None:
                    # Для дистрибьюторов с 1 колонкой записываем лучший имеющийся прайс
                    price_to_put = val_100 or val_25
                    if price_to_put is not None:
                        _set_price_cell(sheet, matched_row, col_25, price_to_put, prev_prices)
                else:
                    if val_25 is not None:
                        _set_price_cell(sheet, matched_row, col_25, val_25, prev_prices)
                    if val_100 is not None:
                        _set_price_cell(sheet, matched_row, col_100, val_100, prev_prices)
                            
        except Exception as ex:
            import logging
            logging.getLogger(__name__).exception(f"Error filling template from file {path.name}: {ex}")
            if progress:
                progress(f"⚠️ Пропущен: Ошибка при обработке {path.name}: {str(ex)}")

    sheet.cell(row=3, column=1).value = f"{today} г."
    wb.save(out)
    return out


def run_svodny_vertical(
    settings_dict: Dict[str, Any],
    catalog_list: List[Tuple[str, List[str]]],
    price_files: List[Dict[str, Any]],
    upload_dir: Path,
    output_dir: Path,
    progress: Optional[Callable[[str], None]] = None,
) -> Path:
    extracted, catalog, labels = _extract_matrix_rows(
        settings_dict=settings_dict,
        catalog_list=catalog_list,
        price_files=price_files,
        upload_dir=upload_dir,
        progress=progress,
    )
    matrix = build_vertical_matrix(extracted, catalog, labels)
    today = datetime.now().strftime("%d.%m.%Y")
    out = output_dir / f"Сводный прайс вертикально {today}.xlsx"
    return _save_matrix(matrix, out, vertical=True)


def run_assortment(
    settings_dict: Dict[str, Any],
    price_files: List[Dict[str, Any]],
    upload_dir: Path,
    output_dir: Path,
    progress: Optional[Callable[[str], None]] = None,
) -> Path:
    key_c, name_c, price_c, stock_c, _ = _column_lists(settings_dict)
    
    # Filter files that actually exist
    valid_entries = []
    for e in price_files:
        filepath = upload_dir / e["filename"]
        if filepath.is_file():
            valid_entries.append(e)

    if not valid_entries:
        raise ValueError("Нет файлов прайсов на сервере. Пожалуйста, загрузите прайсы.")

    rows: List[dict] = []
    total = len(valid_entries)
    for i, e in enumerate(valid_entries, 1):
        path = upload_dir / e["filename"]
        if progress:
            progress(f"Обработка [{i}/{total}]: {path.name}")
        preset = (e.get("read_preset") or "").strip() or None
        disp = _display_label(e.get("display_name", ""), e["filename"])
        try:
            rows.append(
                assortment.count_assortment_for_file(
                    path,
                    key_candidates=key_c,
                    name_candidates=name_c,
                    price_candidates=price_c,
                    stock_candidates=stock_c,
                    read_preset=preset,
                    source_display=disp,
                )
            )
        except Exception as ex:
            rows.append(
                {
                    "file": path.name,
                    "source": disp,
                    "assortment": None,
                    "rows_total": None,
                    "rows_non_empty_id": None,
                    "id_column": None,
                    "id_kind": None,
                    "error": str(ex),
                }
            )

    out_df = pd.DataFrame(rows)
    if "error" in out_df.columns:
        out_df["_ok"] = out_df["error"].isna()
        out_df = out_df.sort_values(
            by=["_ok", "assortment", "source"], ascending=[False, False, True]
        ).drop(columns=["_ok"])
        
    today = datetime.now().strftime("%d.%m.%Y")
    out = output_dir / f"Ассортимент товаров ({today}).xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    out_df.to_excel(out, index=False, engine="openpyxl")
    wb = openpyxl.load_workbook(out)
    for sheet in wb.worksheets:
        assortment.format_summary_sheet(sheet)
    wb.save(out)
    return out


def run_cheapest(output_dir: Path) -> str:
    """Подсвечивает минимальные цены во всех подходящих XLSX в папке результатов."""
    return cheapest.highlight_minimum_prices(output_dir)


def run_best_suppliers_analysis(
    output_dir: Path,
    progress: Optional[Callable[[str], None]] = None,
) -> Path:
    """Анализ последнего сводного прайса: лучшие оптовики и сводка «корзина»."""
    return best_suppliers.run_analysis_to_excel(
        inputs_dir=output_dir,
        report_directory=output_dir,
        progress=progress,
        emit_console_report=False,
    )
