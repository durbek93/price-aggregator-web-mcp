# -*- coding: utf-8 -*-
from __future__ import annotations

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def format_summary_sheet(sheet: openpyxl.worksheet.worksheet.Worksheet) -> None:
    """Приводит лист к аккуратному читаемому виду."""
    if sheet.max_row == 0 or sheet.max_column == 0:
        return

    def pixels_to_width(pixels: int) -> float:
        return round(max((pixels - 5) / 7, 0), 2)

    first_col_width = pixels_to_width(265)
    other_cols_width = pixels_to_width(150)
    header_fill = PatternFill(fill_type="solid", fgColor="D9E1F2")
    thin_side = Side(style="thin", color="D9D9D9")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    for col_idx in range(1, sheet.max_column + 1):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        sheet.column_dimensions[col_letter].width = first_col_width if col_idx == 1 else other_cols_width

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 28

    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=sheet.max_column):
        for cell in row:
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            if cell.column == 1 and 2 <= cell.row <= 99:
                cell.alignment = Alignment(
                    horizontal="left",
                    vertical="center",
                    wrap_text=True,
                )
            cell.border = thin_border
            if cell.row == 1:
                cell.font = Font(bold=True)
                cell.fill = header_fill
            elif cell.column >= 2 and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"
